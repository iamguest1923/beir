"""Build the requested Vietnamese Word report from audited, scoped artifacts."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MODELS = ["BM25", "ANCE", "TAS-B", "DPR", "DeepCT", "SPARTA", "BM25+CE", "ColBERT"]
DATASETS = ["nfcorpus", "scifact", "arguana", "scidocs", "fiqa", "trec-covid",
            "uit-viquad", "vihealthqa", "zalo-legal"]
LABELS = {"nfcorpus": "NFCorpus", "scifact": "SciFact", "arguana": "ArguAna",
          "scidocs": "SCIDOCS", "fiqa": "FiQA", "trec-covid": "TREC-COVID",
          "uit-viquad": "UIT-ViQuAD", "vihealthqa": "ViHealthQA", "zalo-legal": "Zalo Legal"}
NAVY = "17365D"
TEAL = "087E8B"
GRAY = "536273"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def save(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def fmt(value, decimals=5):
    return f"{float(value):.{decimals}f}".replace(".", ",")


def integer(value):
    return f"{value:,}".replace(",", ".")


def shade(cell, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def hyperlink(paragraph, title, url):
    rel = paragraph.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def paragraph(doc, text="", style=None):
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.widow_control = True
    return p


def heading(doc, text):
    doc.add_heading(text, level=1)


def note(doc, text):
    p = paragraph(doc, text)
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(GRAY)
    return p


def table(doc, headers, rows, widths, font=10, best=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = str(text)
        shade(t.rows[0].cells[i], NAVY)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, data in enumerate(rows):
        cells = t.add_row().cells
        for col_index, value in enumerate(data):
            cells[col_index].text = str(value)
            if row_index % 2 == 0:
                shade(cells[col_index], "F0F5F9")
            if best and (row_index, col_index) in best:
                shade(cells[col_index], "DCEEE8")
                for r in cells[col_index].paragraphs[0].runs:
                    r.bold = True
    for row_index, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                m = OxmlElement("w:" + edge)
                m.set(qn("w:w"), "65")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                if col_index:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(font)
                    if row_index == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
    paragraph(doc)
    return t


def section(doc, landscape=False):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    s.top_margin = Cm(1.6 if landscape else 2.0)
    s.bottom_margin = Cm(1.6 if landscape else 2.0)
    s.left_margin = s.right_margin = Cm(1.6 if landscape else 2.2)
    return s


def plot_heatmap(values, output, metric, datasets=None):
    datasets = DATASETS if datasets is None else datasets
    matrix = np.array([[values[d][m] for m in MODELS] for d in datasets])
    fig, ax = plt.subplots(figsize=(11.8, max(1.8, len(datasets) * .3 + .55)))
    im = ax.imshow(matrix, cmap="YlGnBu", vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(range(len(MODELS)), MODELS, fontsize=9)
    ax.set_yticks(range(len(datasets)), [LABELS[d] for d in datasets], fontsize=9)
    ax.tick_params(length=0)
    for i in range(len(datasets)):
        for j in range(len(MODELS)):
            ax.text(j, i, f"{matrix[i,j]:.3f}", ha="center", va="center", fontsize=8,
                    color="white" if matrix[i,j] > .65 else "#17365D")
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.colorbar(im, ax=ax, fraction=.025, pad=.025, label=metric)
    fig.tight_layout(pad=.6)
    fig.savefig(output, dpi=240, bbox_inches="tight", facecolor="white")
    fig.savefig(output.with_suffix(".pdf"), bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build(source, output):
    output.mkdir(parents=True, exist_ok=True)
    with (source / "coverage.csv").open(encoding="utf-8-sig") as stream:
        records = [r for r in csv.DictReader(stream) if r["model"] in MODELS and r["dataset"] in DATASETS]
    assert len(records) == 72 and all(r["status"] == "SUCCESS" for r in records)
    assert len({(r["model"], r["dataset"]) for r in records}) == 72
    metrics = {(r["model"], r["dataset"]): read(Path(r["directory"]) / "metrics.json") for r in records}
    ndcg = {d: {m: metrics[m, d]["ndcg"]["NDCG@10"] for m in MODELS} for d in DATASETS}
    recall = {d: {m: metrics[m, d]["recall"]["Recall@100"] for m in MODELS} for d in DATASETS}
    snapshot = read(source / "snapshot.json")
    captured = datetime.fromisoformat(snapshot["captured_at"]).strftime("%H:%M ngày %d/%m/%Y (UTC+7)")
    inputs = read(source / "inputs.json")
    # Copy only the requested scope into the report's supporting artifacts.
    for filename, values in [("ndcg_at_10.csv", ndcg), ("recall_at_100.csv", recall)]:
        with (output / filename).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["dataset", *MODELS])
            writer.writerows([[d, *[values[d][m] for m in MODELS]] for d in DATASETS])
    with (output / "sources.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    dirs = {r["directory"] for r in records}
    save(output / "verified_artifacts.json", {p: v for p, v in read(source / "verified_artifacts.json").items() if p in dirs})
    save(output / "inputs.json", {d: inputs[d] for d in DATASETS})
    save(output / "report_snapshot.json", {"captured_at": snapshot["captured_at"], "datasets": DATASETS,
         "models": MODELS, "verified_pairs": 72, "expected_pairs": 72,
         "source_audit": str(source.resolve())})
    plot_heatmap(ndcg, output / "ndcg_heatmap.png", "nDCG@10")
    plot_heatmap(recall, output / "recall_heatmap.png", "Recall@100")
    means_en = {m: sum(ndcg[d][m] for d in DATASETS[:6]) / 6 for m in MODELS}
    means_vi = {m: sum(ndcg[d][m] for d in DATASETS[6:]) / 3 for m in MODELS}
    fig, ax = plt.subplots(figsize=(8.0, 3.6))
    x = np.arange(len(MODELS))
    ax.bar(x-.18, [means_en[m] for m in MODELS], .36, label="6 bộ tiếng Anh", color="#17365D")
    ax.bar(x+.18, [means_vi[m] for m in MODELS], .36, label="3 bộ tiếng Việt", color="#13A3A5")
    ax.set_xticks(x, MODELS, fontsize=9)
    ax.set_ylim(0, .7)
    ax.set_ylabel("nDCG@10 trung bình theo dataset")
    ax.grid(axis="y", alpha=.2)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, ncol=2)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    for extension in ["png", "pdf"]:
        fig.savefig(output / f"macro_comparison.{extension}", dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
    s.header_distance = s.footer_distance = Cm(.8)
    styles = doc.styles
    for name in ["Normal", "Body Text", "List Bullet", "List Number"]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(12)
        styles[name].paragraph_format.space_after = Pt(5)
        styles[name].paragraph_format.line_spacing = 1.12
    for name in ["Heading 1", "Heading 2", "Title", "Subtitle"]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 1"].paragraph_format.space_before = Pt(3)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "vi-VN")
    styles["Normal"].element.get_or_add_rPr().append(lang)
    header = s.header.paragraphs[0]
    header.text = "BEIR  |  ĐÁNH GIÁ TRUY HỒI THÔNG TIN"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(GRAY)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Báo cáo thực nghiệm • ")
    for instruction in ["PAGE", "NUMPAGES"]:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        footer._p.append(field)
        if instruction == "PAGE":
            footer.add_run(" / ")
    for r in footer.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.core_properties.title = "Báo cáo kết quả đánh giá 9 dataset với 8 mô hình retrieval"
    doc.core_properties.subject = "72 kết quả thực nghiệm đã kiểm tra từ artifact BEIR"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "BEIR, retrieval, BM25, tiếng Việt, nDCG@10, Recall@100"

    # 1. Cover and executive summary.
    paragraph(doc)
    p = paragraph(doc, "BÁO CÁO KẾT QUẢ\nĐÁNH GIÁ", "Title")
    for r in p.runs:
        r.font.size = Pt(30)
        r.bold = True
    paragraph(doc, "TRUY HỒI THÔNG TIN TRÊN BEIR\nVÀ CÁC BỘ DỮ LIỆU TIẾNG VIỆT", "Subtitle")
    paragraph(doc)
    table(doc, ["DATASET", "MÔ HÌNH", "KẾT QUẢ HỢP LỆ"], [["9", "8", "72 / 72"]], [5.5, 5.5, 5.5], font=15)
    paragraph(doc, "Phạm vi mô hình: BM25, ANCE, TAS-B, DPR, DeepCT, SPARTA, BM25+CE và ColBERT.")
    paragraph(doc, "Phạm vi dữ liệu: NFCorpus, SciFact, ArguAna, SCIDOCS, FiQA, TREC-COVID, UIT-ViQuAD, ViHealthQA và Zalo Legal.")
    paragraph(doc, "Thời điểm chốt số liệu: " + captured + ".")
    paragraph(doc, "Tài liệu tổng hợp các kết quả chạy thực tế trong project; không điền điểm từ bài báo vào bảng thực nghiệm.")
    doc.add_heading("Kết quả nổi bật", level=2)
    for text in [
        "BM25+CE đạt nDCG@10 cao hơn BM25 trên cả sáu bộ BEIR tiếng Anh được khảo sát.",
        "BM25 đạt nDCG@10 cao nhất trên cả ba bộ tiếng Việt trong phạm vi tám mô hình đã chạy.",
        "Kết quả phụ thuộc mạnh vào dataset và ngôn ngữ; không có mô hình đứng đầu trên toàn bộ chín bộ dữ liệu.",
    ]:
        paragraph(doc, text, "List Bullet")
    note(doc, "Các kết luận chỉ áp dụng cho checkpoint, dữ liệu và cấu hình được ghi trong báo cáo; chưa có kiểm định ý nghĩa thống kê.")

    # 2. Dataset scope.
    doc.add_page_break()
    heading(doc, "1. Mục tiêu và phạm vi dữ liệu")
    paragraph(doc, "Mục tiêu là xây dựng quy trình thống nhất để đánh giá khả năng tìm tài liệu liên quan của tám mô hình trên chín dataset. Sáu bộ BEIR tiếng Anh được dùng làm nền so sánh; ba bộ tiếng Việt mở rộng đánh giá sang kiến thức tổng quát, sức khỏe và pháp luật. Toàn bộ kết quả trong tài liệu dùng split test.")
    domains = ["Y sinh / dinh dưỡng", "Kiểm chứng khoa học", "Truy hồi phản biện", "Tài liệu khoa học / trích dẫn", "Tài chính", "Nghiên cứu COVID-19", "Wikipedia / kiến thức chung", "Sức khỏe / tư vấn y tế", "Pháp luật Việt Nam"]
    data_rows = []
    for index, d in enumerate(DATASETS):
        with (Path(inputs[d]["path"]) / "qrels/test.tsv").open(encoding="utf-8") as stream:
            rel_count = sum(1 for _ in csv.DictReader(stream, delimiter="\t"))
        data_rows.append([LABELS[d], "Việt" if index >= 6 else "Anh", domains[index],
                          integer(inputs[d]["documents"]), integer(inputs[d]["queries"]), integer(rel_count)])
    table(doc, ["Dataset", "Ngôn ngữ", "Miền / tác vụ", "Document", "Query test", "Qrels test"],
          data_rows, [2.55, 1.3, 4.2, 2.0, 2.0, 2.55], font=10)
    total_docs = sum(inputs[d]["documents"] for d in DATASETS)
    total_queries = sum(inputs[d]["queries"] for d in DATASETS)
    paragraph(doc, f"Tổng theo từng bộ: {integer(total_docs)} document và {integer(total_queries)} query test. Qrels là các cặp query–document có nhãn liên quan; một query có thể có nhiều document liên quan.")
    paragraph(doc, "NFCorpus và TREC-COVID có nhiều tài liệu liên quan cho một query; do đó Recall@100 có thể thấp dù các vị trí đầu được xếp hạng tốt. Không so sánh trực tiếp độ khó của hai dataset chỉ bằng một giá trị Recall@100.")
    note(doc, "Corpus của mỗi bộ được giữ riêng. Các số đếm là từ file local đã dùng để đánh giá, không phải ước lượng từ mô tả trên website.")

    # 3. Preparation policy.
    doc.add_page_break()
    heading(doc, "2. Chuẩn bị ba dataset tiếng Việt")
    paragraph(doc, "Dữ liệu được chuẩn hóa về định dạng BEIR: corpus.jsonl chứa document; queries.jsonl chứa query; qrels/test.tsv chứa query-id, corpus-id và score. Script chuẩn bị nằm trong beir/new/data/prepare.py. Metadata lưu nguồn tải, commit, checksum, số lượng và thống kê xử lý.")
    table(doc, ["Dataset", "Nguồn và phép chuyển đổi"], [
        ["UIT-ViQuAD", "Mirror vile99/vi_quad; question → query, context → document. Chỉ giữ query có đáp án và kiểm tra đáp án nằm trong context. Context vẫn được giữ khi query không trả lời được. [3]"],
        ["ViHealthQA", "Nguồn tarudesu/ViHealthQA; question → query, answer passage → document. Không đưa chính câu hỏi vào tiêu đề document. Đây là bản chuyển đổi QA sang retrieval. [4]"],
        ["Zalo Legal", "Nguồn GreenNode/zalo-ai-legal-text-retrieval-vn; giữ ID điều luật, query ID, split và các positive gốc của bản retrieval. [5]"],
    ], [2.7, 13.9], font=11)
    paragraph(doc, "Với dữ liệu QA, passage được chuẩn hóa Unicode NFC và whitespace; passage trùng được gộp theo hash nội dung. Corpus chứa passage của các split nguồn; các query test chỉ được dùng để đánh giá. Query QA có prefix split nhằm tránh va chạm ID cục bộ.")
    table(doc, ["Dataset", "Train query", "Dev query", "Test query", "Chính sách cần lưu ý"], [
        ["UIT-ViQuAD", "15.333", "1.937", "1.970", "Bỏ 9.217 query unanswerable; bản mirror không được khẳng định là bản 1.1 gốc."],
        ["ViHealthQA", "7.009", "993", "2.013", "Giữ theo tên file nguồn; val/test khác phần mô tả trong dataset card."],
        ["Zalo Legal", "2.432", "—", "788", "Giữ 357 document chỉ có title; gộp 102 dòng query ID trùng cùng nội dung."],
    ], [2.55, 1.75, 1.65, 1.75, 8.9], font=10)
    paragraph(doc, "Kiểm tra overlap phát hiện query trùng văn bản giữa train và test: UIT-ViQuAD có 1, ViHealthQA có 9 và Zalo Legal có 24. Riêng Zalo Legal, 24 query ID xuất hiện ở cả hai split. Cần kiểm soát overlap trước mọi thí nghiệm supervised fine-tuning. Các lượt đánh giá trong báo cáo không fine-tune tám model bằng nhãn train của các bộ này.")
    note(doc, "Qrels chuyển từ QA chỉ xác nhận passage đi kèm câu hỏi; các passage khác vẫn có thể liên quan nhưng chưa được gán nhãn. Điểm retrieval không thay thế điểm extractive/generative QA.")

    # 4. Methods.
    doc.add_page_break()
    heading(doc, "3. Các mô hình được đánh giá")
    paragraph(doc, "Các checkpoint và pipeline được giữ nhất quán giữa các dataset để quan sát khả năng tổng quát hóa. Bảng dưới mô tả cấu hình đã triển khai; việc dùng checkpoint cùng họ với bài báo không đồng nghĩa tái lập chính xác toàn bộ môi trường gốc. [1, 2]")
    table(doc, ["Model", "Cấu hình / checkpoint", "Cách truy hồi"], [
        ["BM25", "Anserini qua Pyserini 0.12.0; k1 = 0,9; b = 0,4; contents/title có trọng số 1,0.", "Khớp từ và xếp hạng lexical trên corpus."],
        ["ANCE", "sentence-transformers/msmarco-roberta-base-ance-firstp", "Bi-encoder; exact search với dot product."],
        ["TAS-B", "sentence-transformers/msmarco-distilbert-base-tas-b", "Bi-encoder; exact search với dot product."],
        ["DPR", "sentence-transformers/facebook-dpr-question_encoder-multiset-base và facebook-dpr-ctx_encoder-multiset-base", "Hai encoder riêng cho query/context; title [SEP] text; exact dot product."],
        ["DeepCT", "macavaney/deepct; checkpoint gốc được chuyển sang PyTorch; scaling = 100, keep_all_terms = True.", "Dự đoán trọng số từ, tạo pseudo-document và tìm bằng Anserini BM25."],
        ["SPARTA", "BeIR/sparta-msmarco-distilbert-base-v1; giữ 2.000 phần tử khác 0 trong vector sparse.", "Biểu diễn sparse học được và SparseSearch."],
        ["BM25+CE", "cross-encoder/ms-marco-MiniLM-L-6-v2; candidate_k = 100; batch CE = 8.", "Lấy top-100 từ BM25, sau đó xếp lại theo điểm Cross-Encoder."],
        ["ColBERT", "Checkpoint ColBERT v1 msmarco.psg.l2; doc_maxlen = 300; query_maxlen = 32.", "IVFPQ ANN với FAISS, sau đó MaxSim; nprobe = 32, trả về 100 kết quả."],
    ], [1.8, 8.4, 6.4], font=10)
    paragraph(doc, "MiniLM-L6 được chọn để bám baseline BM25+CE: mục 4 của bài báo BEIR mô tả reranking top-100 bằng MiniLM 6 tầng, hidden size 384; bảng 6 nêu đúng checkpoint ms-marco-MiniLM-L-6-v2. Model được huấn luyện trên MS MARCO tiếng Anh. [1, 6]")
    note(doc, "Các mô hình pretrained được áp dụng trực tiếp cho tiếng Việt; chưa thay bằng checkpoint đa ngôn ngữ hoặc checkpoint chuyên biệt tiếng Việt.")

    # 5. Hardware and metrics.
    doc.add_page_break()
    heading(doc, "4. Môi trường và quy trình đánh giá")
    table(doc, ["Thành phần", "Cấu hình đã sử dụng"], [
        ["CPU", "2 × Intel Xeon Silver 4210R @ 2,40 GHz; tổng 20 nhân / 40 luồng."],
        ["GPU", "1 × NVIDIA A30, 24 GiB VRAM."],
        ["RAM", "Khoảng 128 GB; hệ điều hành nhận 125,5 GiB."],
        ["Môi trường chính", "Python 3.9; BEIR 2.2.0; PyTorch 2.8.0; Transformers 4.57.6; Sentence Transformers 5.1.2; Pyserini 0.12.0; pytrec-eval-terrier 0.5.10."],
    ], [3.4, 13.2], font=11)
    paragraph(doc, "ColBERT dùng runtime tách riêng để tương thích mã v1; không suy diễn phiên bản thư viện chính cho runtime này. Config và metadata của từng lượt được giữ lại để kiểm tra. Hàng đợi bổ sung chỉ chạy một model GPU tại một thời điểm; BM25 có lane CPU riêng. Máy còn có tiến trình chạy riêng, nên tài liệu này không dùng thời gian wall-clock để kết luận model nào nhanh hơn.")
    paragraph(doc, "Dense retrieval mã hóa query và document theo batch; corpus được chia thành chunk tối đa 10.000 document. ANCE dùng batch 16; TAS-B/DPR dùng batch 32. Giới hạn đầu vào phổ biến là 512 token; ColBERT dùng cấu hình độ dài riêng. Các neural encoder trong pipeline chính dùng FP32.")
    doc.add_heading("Chỉ số đánh giá", level=2)
    paragraph(doc, "nDCG@10 là chỉ số chính: đo chất lượng thứ tự của 10 document đầu, có xét mức độ liên quan trong qrels. Điểm được chuẩn hóa theo thứ tự lý tưởng của từng query; giá trị càng lớn càng tốt. Khi đọc bảng, cần so sánh các model trên cùng một dataset.")
    paragraph(doc, "Recall@100 là tỷ lệ document liên quan xuất hiện trong 100 kết quả đầu, tính theo từng query rồi lấy trung bình. Ví dụ, một query có rất nhiều positive thì Recall@100 bị giới hạn bởi độ sâu truy hồi dù các kết quả đầu đều tốt.")
    paragraph(doc, "Code còn lưu MAP, Precision và MRR ở các mức k đã cấu hình. Báo cáo tập trung nDCG@10 và Recall@100 để bảng gọn và so sánh nhất quán. Khi query ID trùng document ID, pipeline loại self-match theo quy ước BEIR.")
    note(doc, "Seed 42 được ghi trong cấu hình dùng chung, nhưng không phải mọi luồng inference đều gọi hàm đặt seed. Báo cáo không tuyên bố tính xác định tuyệt đối giữa các phần cứng/môi trường.")

    # 6. Landscape matrices.
    section(doc, True)
    heading(doc, "5. Kết quả nDCG@10")
    paragraph(doc, "Đủ 72 kết quả hợp lệ. Ô tô xanh và chữ đậm đánh dấu điểm cao nhất trong tám model của từng dataset. Điểm trình bày với 5 chữ số thập phân; dữ liệu gốc lưu trong CSV đi kèm.")
    best = {(i, j+1) for i, d in enumerate(DATASETS) for j, m in enumerate(MODELS)
            if ndcg[d][m] == max(ndcg[d].values())}
    table(doc, ["Dataset", *MODELS], [[LABELS[d], *[fmt(ndcg[d][m]) for m in MODELS]] for d in DATASETS],
          [3.15] + [2.9]*8, font=10, best=best)
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(output / "ndcg_heatmap.png"), width=Cm(23.5))
    note(doc, "Hình 1. Phân bố nDCG@10 trên chín dataset. Thang màu cố định 0–1; hình chỉ minh họa các số đã có trong bảng.")

    doc.add_page_break()
    heading(doc, "6. Kết quả Recall@100")
    paragraph(doc, "Recall@100 phản ánh độ bao phủ nhãn liên quan trong top-100. BM25+CE giữ cùng tập ứng viên top-100 của BM25 nên Recall@100 bằng nhau, dù thứ tự và nDCG@10 có thể thay đổi.")
    best = {(i, j+1) for i, d in enumerate(DATASETS) for j, m in enumerate(MODELS)
            if recall[d][m] == max(recall[d].values())}
    table(doc, ["Dataset", *MODELS], [[LABELS[d], *[fmt(recall[d][m]) for m in MODELS]] for d in DATASETS],
          [3.15] + [2.9]*8, font=10, best=best)
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(output / "recall_heatmap.png"), width=Cm(23.5))
    note(doc, "Hình 2. Recall@100 trên cùng dữ liệu. Điểm TREC-COVID thấp về Recall@100 cần được đọc cùng số lượng positive lớn; không thể diễn giải đơn giản thành lỗi pipeline.")

    # 8. Fair macro averages.
    section(doc)
    heading(doc, "7. Phân tích kết quả tổng hợp")
    paragraph(doc, "Để tránh so sánh trên các tập con khác nhau, bảng dưới tính trung bình nDCG@10 trên cùng sáu bộ tiếng Anh và cùng ba bộ tiếng Việt cho tất cả tám model. Mỗi dataset có trọng số bằng nhau; đây là macro-average theo dataset, không phải trung bình theo số query.")
    table(doc, ["Model", "TB 6 bộ tiếng Anh", "TB 3 bộ tiếng Việt"],
          [[m, fmt(means_en[m]), fmt(means_vi[m])] for m in MODELS], [4.2, 6.2, 6.2], font=11)
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(output / "macro_comparison.png"), width=Cm(15.8))
    note(doc, "Hình 3. So sánh trung bình trong hai nhóm dataset khác nhau. Không quy toàn bộ chênh lệch giữa hai nhóm cho ngôn ngữ vì miền và tác vụ cũng khác nhau.")
    paragraph(doc, f"Trên nhóm tiếng Anh, BM25+CE đạt trung bình {fmt(means_en['BM25+CE'])}, cao hơn BM25 ({fmt(means_en['BM25'])}) và cải thiện trên cả sáu dataset. ArguAna là ngoại lệ về model đứng đầu: DeepCT đạt {fmt(ndcg['arguana']['DeepCT'])}, nhỉnh hơn TAS-B ({fmt(ndcg['arguana']['TAS-B'])}) và BM25+CE ({fmt(ndcg['arguana']['BM25+CE'])}).")
    paragraph(doc, f"Trên nhóm tiếng Việt, BM25 đạt trung bình {fmt(means_vi['BM25'])}, cao nhất trong tám model. BM25 cũng đứng đầu riêng từng bộ. BM25+CE đứng thứ hai cả ba bộ; ColBERT có trung bình {fmt(means_vi['ColBERT'])}. Các điểm số này cho thấy thứ hạng model trên tiếng Anh không bảo đảm được giữ nguyên khi chuyển sang tiếng Việt.")

    # 9. CE diagnostics.
    doc.add_page_break()
    heading(doc, "8. Vì sao BM25+CE giảm điểm trên tiếng Việt?")
    table(doc, ["Dataset", "BM25", "BM25+CE", "Δ nDCG@10", "Giảm tương đối"], [
        [LABELS[d], fmt(ndcg[d]["BM25"]), fmt(ndcg[d]["BM25+CE"]),
         fmt(ndcg[d]["BM25+CE"]-ndcg[d]["BM25"]),
         fmt((ndcg[d]["BM25"]-ndcg[d]["BM25+CE"])/ndcg[d]["BM25"]*100, 1)+"%"]
        for d in DATASETS[6:]], [3.0, 2.7, 3.0, 3.5, 2.4], font=11)
    paragraph(doc, "Kiểm tra trực tiếp các runfile cho thấy top-100 ứng viên của BM25 và BM25+CE giống nhau ở toàn bộ 1.970 query ViQuAD, 2.013 query ViHealthQA và 788 query Zalo Legal. Recall@100 tương ứng giữ nguyên. Vì vậy, mức giảm nDCG@10 trong các lượt này xuất hiện ở bước đổi thứ tự ứng viên, không phải do thay corpus hoặc mất ứng viên trước reranking.")
    paragraph(doc, "Cross-Encoder chấm từng cặp query–document, sau đó pipeline sắp xếp hoàn toàn theo điểm CE. Điểm BM25 không được trộn lại để giữ thứ tự ban đầu. Nếu CE đánh giá sai mức liên quan, positive có thể bị đẩy xuống vị trí thấp hơn trong cùng tập ứng viên.")
    paragraph(doc, "Checkpoint MiniLM đang dùng được công bố cho tiếng Anh, huấn luyện trên MS MARCO. Việc áp dụng trực tiếp cho câu hỏi, đoạn y tế và điều luật tiếng Việt tạo ra chênh lệch ngôn ngữ và miền dữ liệu. Đây là lời giải thích phù hợp với mẫu kết quả: CE cải thiện trên sáu bộ tiếng Anh nhưng giảm trên cả ba bộ tiếng Việt. [6]")
    paragraph(doc, "Tuy nhiên, thí nghiệm hiện tại chưa tách riêng ảnh hưởng của ngôn ngữ, miền, độ dài và tokenizer. Do đó báo cáo coi đây là giả thuyết được dữ liệu hỗ trợ, không khẳng định quan hệ nhân quả đã được chứng minh. Một kiểm chứng tiếp theo phù hợp là giữ nguyên top-100 và đánh giá một reranker đa ngôn ngữ/tiếng Việt ở cấu hình riêng.")
    paragraph(doc, "Việc giữ MiniLM-L6 trong thí nghiệm này có mục đích đo khả năng chuyển giao của baseline BEIR. Không nên mô tả kết quả hiện tại là hiệu năng tối ưu của mọi phương pháp Cross-Encoder trên tiếng Việt.")

    # 10. Validation and limitations.
    doc.add_page_break()
    heading(doc, "9. Kiểm tra tính đúng và giới hạn")
    paragraph(doc, "Mỗi kết quả được đối chiếu status SUCCESS, dataset/split, số document/query, checksum đầu vào nếu có, checksum metric/ranking, ID trong runfile và các metric tính lại từ ranking. Bản báo cáo chỉ chọn 72 cặp đã qua audit; không lấy các dòng FAILED từ CSV cũ làm kết quả cuối cùng.")
    paragraph(doc, "TREC runfile không có dòng cho query không tìm được tài liệu. Audit khôi phục các query này với ranking rỗng trước khi tính lại metric, để không bỏ query khỏi mẫu số. Với các phiên BM25/ANCE cũ chưa lưu input hash, audit kiểm tra đường dẫn dữ liệu đã ghi, file hiện có, số lượng và metric tính lại; vẫn không thể chứng minh dữ liệu nguồn trong quá khứ chưa từng thay đổi.")
    table(doc, ["Vấn đề đã gặp", "Xử lý và kết quả"], [
        ["BM25 thiếu document trong index", "Anserini mặc định bỏ contents rỗng. Đã dùng generator giữ document gốc, kể cả title-only, và kiểm tra đủ số document sau index."],
        ["Ba corpus BM25 từng lỗi", "SCIDOCS: 344; FiQA: 38; TREC-COVID: 42.140 document bị bỏ ở phiên cũ. Các lượt chạy bổ sung đã thành công và được audit."],
        ["Bộ nhớ khi ColBERT MaxSim", "Chấm theo nhóm 256 candidates thay vì tạo bộ đệm lớn; giữ checkpoint, thuật toán và tập candidate của pipeline."],
    ], [5.0, 11.6], font=10.5)
    doc.add_heading("Giới hạn khi diễn giải", level=2)
    for text in [
        "Qrels QA chuyển đổi có thể chưa đầy đủ; corpus chứa passage từ mọi split và cần kiểm soát query overlap trước huấn luyện supervised.",
        "ArguAna có 5 document ID được qrels tham chiếu nhưng không có trong corpus local. Giữ query/qrels gốc; các positive thiếu không thể được truy hồi và vẫn ảnh hưởng mẫu số liên quan.",
        "Các neural model chủ yếu được học từ dữ liệu tiếng Anh; chưa tối ưu cho ba bộ tiếng Việt. Bản ViQuAD dùng mirror có unanswerable, không khẳng định tái lập bản 1.1 gốc.",
        "Chưa chạy nhiều seed hoặc kiểm định thống kê theo query. Chênh lệch nhỏ giữa các model cần được xem là mô tả của lượt chạy này.",
        "Không dùng số liệu trong báo cáo để khẳng định trùng điểm bài báo BEIR hoặc model luôn tốt hơn trên mọi dữ liệu.",
    ]:
        paragraph(doc, text, "List Bullet")

    # 11. Conclusions and sources.
    doc.add_page_break()
    heading(doc, "10. Kết luận và khả năng tái lập")
    paragraph(doc, "Đã hoàn thành đủ 72 thí nghiệm trong phạm vi chín dataset và tám model. Quy trình chuẩn hóa dữ liệu, chạy mô hình, ghi ranking, chấm metric và kiểm tra artifact đã được thống nhất. Các kết quả cho thấy BM25 là baseline mạnh trên ba bộ tiếng Việt, còn BM25+CE đem lại lợi ích nhất quán so với BM25 trên sáu bộ BEIR tiếng Anh được khảo sát.")
    paragraph(doc, "Hướng nghiên cứu tiếp theo là đánh giá checkpoint đa ngôn ngữ hoặc tiếng Việt, kiểm soát overlap, bổ sung nhãn liên quan cho dữ liệu QA chuyển đổi, và kiểm định chênh lệch trên cùng tập query. Mỗi cấu hình mới cần giữ riêng kết quả để không thay thế hoặc trộn với baseline trong báo cáo này.")
    doc.add_heading("Artifact đi kèm", level=2)
    table(doc, ["File", "Nội dung"], [
        ["ndcg_at_10.csv / recall_at_100.csv", "Hai bảng 9 dataset × 8 model; giữ giá trị gốc."],
        ["sources.csv", "72 dòng kết quả; đường dẫn phiên chạy gốc và trạng thái."],
        ["inputs.json / verified_artifacts.json", "Số lượng, hash dữ liệu và kết quả kiểm tra ranking/metric."],
        ["report_snapshot.json", "Phạm vi, thời điểm chốt và nguồn audit của báo cáo."],
        ["Các hình .png / .pdf", "Biểu đồ độc lập, có thể dùng lại trong bài viết hoặc slide."],
    ], [6.2, 10.4], font=10)
    paragraph(doc, "Mã chạy: beir/runs/<model>/evaluate_*.py. Mã bổ sung và audit: beir/runs/complete_benchmark.py. Ranking và metadata gốc được lưu ở beir/runs/<model>/results/<run_id>/<dataset>/. File sources.csv chỉ ra chính xác phiên nào được dùng cho từng ô kết quả.")
    note(doc, "Các bảng trong file Word là snapshot tại thời điểm chốt; chúng không tự thay đổi khi project có lượt chạy mới.")
    doc.add_heading("Tài liệu và nguồn dữ liệu", level=2)
    references = [
        ("[1] Thakur và cộng sự (2021). BEIR: A Heterogeneous Benchmark for Zero-shot Evaluation of Information Retrieval Models. Mục 4 và bảng 6.", "https://arxiv.org/html/2104.08663v4"),
        ("[2] BEIR — mã nguồn và framework đánh giá.", "https://github.com/beir-cellar/beir"),
        ("[3] UIT-ViQuAD mirror — vile99/vi_quad.", "https://huggingface.co/datasets/vile99/vi_quad"),
        ("[4] ViHealthQA — tarudesu/ViHealthQA.", "https://huggingface.co/datasets/tarudesu/ViHealthQA"),
        ("[5] Zalo Legal retrieval — GreenNode/zalo-ai-legal-text-retrieval-vn.", "https://huggingface.co/datasets/GreenNode/zalo-ai-legal-text-retrieval-vn"),
        ("[6] Model card — cross-encoder/ms-marco-MiniLM-L6-v2.", "https://huggingface.co/cross-encoder/ms-marco-MiniLM-L6-v2"),
    ]
    for title, url in references:
        p = paragraph(doc)
        hyperlink(p, title, url)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(10)

    path = output / "BAO_CAO_KET_QUA_DANH_GIA.docx"
    doc.save(path)
    save(output / "document_manifest.json", {"file": path.name,
         "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "datasets": DATASETS,
         "models": MODELS, "results": 72, "snapshot": snapshot["captured_at"]})
    print(path.resolve())
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
