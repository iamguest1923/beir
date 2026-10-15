"""Reorganize the audited Word report into paper and extension datasets."""
from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import Patch
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_evaluation_report import (
    DATASETS, MODELS, LABELS, fmt, integer, read, save, paragraph,
    heading, note, table, section, plot_heatmap,
)

GROUPS = [
    ("dataset_bai_bao", DATASETS[:6], "PHẦN I. DATASET TRONG BÀI BÁO BEIR", "I"),
    ("dataset_mo_rong", DATASETS[6:], "PHẦN II. DATASET MỞ RỘNG TIẾNG VIỆT", "II"),
]


def blocks_by_heading(doc):
    result = {"cover": []}
    current = "cover"
    for element in list(doc.element.body):
        if element.tag == qn("w:sectPr") or element.xpath(".//w:sectPr"):
            continue
        if element.tag == qn("w:p"):
            # Copy content but replace the original layout/page breaks ourselves.
            if element.xpath('.//w:br[@w:type="page"]'):
                continue
            styles = element.xpath("./w:pPr/w:pStyle/@w:val")
            if styles == ["Heading1"]:
                current = "".join(element.xpath(".//w:t/text()"))
                result[current] = []
                continue
        result[current].append(deepcopy(element))
    return result


def append_blocks(doc, blocks):
    for element in blocks:
        doc.element.body.insert(-1, deepcopy(element))


def chart_means(values, datasets, output):
    means = {m: sum(values[d][m] for d in datasets) / len(datasets) for m in MODELS}
    fig, ax = plt.subplots(figsize=(8, 3.15))
    y = np.arange(len(MODELS))
    colors = ["#087E8B" if means[m] == max(means.values()) else "#17365D" for m in MODELS]
    ax.barh(y, [means[m] for m in MODELS], color=colors, height=.65)
    ax.set_yticks(y, MODELS, fontsize=9)
    ax.invert_yaxis()
    ax.set_xlim(0, .7)
    ax.set_xlabel("nDCG@10 trung bình theo dataset")
    ax.grid(axis="x", alpha=.2)
    ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    for i, m in enumerate(MODELS):
        ax.text(means[m]+.009, i, f"{means[m]:.5f}", va="center", fontsize=9)
    fig.tight_layout()
    for ext in ["png", "pdf"]:
        fig.savefig(output.with_suffix("."+ext), dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return means


def figure(doc, path, caption, width):
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    note(doc, caption)


def reference_bar_charts(values, datasets, output_dir, metric, part):
    """Three dataset panels per figure, matching the user's reference style."""
    charts = []
    chunks = [datasets[i:i+3] for i in range(0, len(datasets), 3)]
    key = "ndcg" if metric == "nDCG@10" else "recall"
    for chunk_index, names in enumerate(chunks):
        fig, axes = plt.subplots(1, len(names), figsize=(16, 5.7), squeeze=False)
        title = ("Kết quả trên các dataset trong bài báo BEIR" if part == "I"
                 else "Kết quả mở rộng trên các dataset tiếng Việt")
        if metric == "Recall@100":
            title += " — Recall@100"
        if len(chunks) > 1:
            title += f" ({chunk_index+1}/{len(chunks)})"
        fig.suptitle(title, fontsize=20, fontweight="bold", y=.985)
        fig.legend(handles=[Patch(color="#478895", label="BM25"),
                            Patch(color="#EE9634", label="Các mô hình còn lại")],
                   loc="upper center", bbox_to_anchor=(.5, .91), ncol=2,
                   frameon=False, fontsize=12)
        for ax, dataset in zip(axes[0], names):
            numbers = [values[dataset][m] for m in MODELS]
            bars = ax.bar(np.arange(len(MODELS)), numbers, width=.7,
                          color=["#478895"] + ["#EE9634"]*(len(MODELS)-1))
            ax.set_title(LABELS[dataset], loc="left", fontsize=14, fontweight="bold", pad=13)
            ax.set_ylabel(metric, fontsize=12)
            ax.set_xticks(np.arange(len(MODELS)), MODELS, rotation=38, ha="right", fontsize=11)
            # All panels use the same scale; leave space for vertical value labels.
            ax.set_ylim(0, 1.18 if metric == "Recall@100" else 1.0)
            ax.set_yticks(np.arange(0, 1.01, .2))
            ax.tick_params(axis="both", length=0, pad=7, labelsize=11)
            ax.set_axisbelow(True)
            ax.grid(axis="y", color="#DDDDDD", linewidth=.8)
            ax.spines[["top", "right"]].set_visible(False)
            for side in ["left", "bottom"]:
                ax.spines[side].set_color("#BDBDBD")
            ax.bar_label(bars, labels=[f"{n:.5f}" for n in numbers], padding=4,
                         rotation=90, fontsize=11)
        fig.subplots_adjust(left=.055, right=.99, top=.765, bottom=.20, wspace=.24)
        fig.text(.5, .025, "Nguồn: kết quả thực nghiệm đã kiểm tra; nhãn cột hiển thị 5 chữ số thập phân.",
                 ha="center", fontsize=10.5, color="#607D8B")
        path = output_dir / f"{key}_cot_theo_dataset_{chunk_index+1}.png"
        for suffix in [".png", ".pdf"]:
            fig.savefig(path.with_suffix(suffix), dpi=260, facecolor="white")
        plt.close(fig)
        save(path.with_suffix(".json"), {"metric": metric, "datasets": names, "models": MODELS,
             "values": {d: {m: values[d][m] for m in MODELS} for d in names}})
        charts.append((path, names))
    return charts


def build(source, output, add_reference_charts=False):
    output.mkdir(parents=True, exist_ok=True)
    original = source / "BAO_CAO_KET_QUA_DANH_GIA.docx"
    doc = Document(original)
    blocks = blocks_by_heading(doc)
    header_footer_refs = [deepcopy(e) for e in doc.sections[0]._sectPr
                          if e.tag in {qn("w:headerReference"), qn("w:footerReference")}]
    values = {}
    for key, filename in [("ndcg", "ndcg_at_10.csv"), ("recall", "recall_at_100.csv")]:
        with (source / filename).open(encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
        assert {r["dataset"] for r in rows} == set(DATASETS)
        values[key] = {r["dataset"]: {m: float(r[m]) for m in MODELS} for r in rows}
    inputs = read(source / "inputs.json")
    with (source / "sources.csv").open(encoding="utf-8-sig") as stream:
        records = list(csv.DictReader(stream))
    assert len(records) == 72 and all(r["status"] == "SUCCESS" for r in records)

    # Keep source relationships, styles, headers and footers; replace the body.
    for element in list(doc.element.body):
        if element.tag != qn("w:sectPr"):
            doc.element.body.remove(element)
    s = doc.sections[0]
    for ref in header_footer_refs:
        s._sectPr.insert(0, ref)
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
    append_blocks(doc, blocks["cover"])
    # Two primary parts; common configuration appears before Part I.
    paragraph(doc, "Cấu trúc báo cáo: Phần I trình bày riêng sáu dataset trong bài báo BEIR (48 kết quả); Phần II trình bày riêng ba dataset mở rộng tiếng Việt (24 kết quả). Mỗi phần có bảng, biểu đồ và nhận xét độc lập.")
    doc.add_page_break()
    doc.add_heading("Thiết lập chung — Tám mô hình đánh giá", level=2)
    append_blocks(doc, blocks["3. Các mô hình được đánh giá"])
    doc.add_page_break()
    doc.add_heading("Thiết lập chung — Môi trường và chỉ số", level=2)
    append_blocks(doc, blocks["4. Môi trường và quy trình đánh giá"])

    # Common audit is prose; no combined table or chart of the two groups.
    paragraph(doc, "Cả hai phần sử dụng kết quả đã qua kiểm tra status, split, số lượng, checksum và tính lại metric từ ranking. Query không có hit được khôi phục với ranking rỗng để giữ đúng mẫu số. Một số phiên BM25/ANCE cũ không lưu input hash ban đầu; audit đối chiếu file hiện có tại đường dẫn đã ghi và không thể khẳng định file nguồn trong quá khứ chưa từng thay đổi.")

    domains = dict(zip(DATASETS, [
        "Y sinh / dinh dưỡng", "Kiểm chứng khoa học", "Truy hồi phản biện",
        "Tài liệu khoa học / trích dẫn", "Tài chính", "Nghiên cứu COVID-19",
        "Wikipedia / kiến thức chung", "Sức khỏe / tư vấn y tế", "Pháp luật Việt Nam",
    ]))
    section_manifest = {}
    for slug, datasets, title, part in GROUPS:
        section(doc)
        heading(doc, title)
        doc.add_heading(f"{part}.1. Phạm vi và dữ liệu", level=2)
        if part == "I":
            paragraph(doc, "Phần này chỉ đánh giá sáu dataset được chọn từ benchmark trong bài báo BEIR: NFCorpus, SciFact, ArguAna, SCIDOCS, FiQA và TREC-COVID. Đây là tập con của benchmark gốc, không đại diện toàn bộ các dataset của bài báo. Cả tám mô hình đã hoàn tất trên cả sáu bộ, tương ứng 48/48 cặp hợp lệ. [1, 2]")
        else:
            paragraph(doc, "Phần này đánh giá ba dataset tiếng Việt được bổ sung vào project: UIT-ViQuAD, ViHealthQA và Zalo Legal. Ba bộ này không nằm trong benchmark tiếng Anh được khảo sát ở Phần I. Giữ cùng tám checkpoint và quy trình chấm để đo khả năng áp dụng trên dữ liệu mở rộng; đã hoàn tất 24/24 cặp hợp lệ.")
        rows = []
        for d in datasets:
            with (Path(inputs[d]["path"]) / "qrels/test.tsv").open(encoding="utf-8") as stream:
                count = sum(1 for _ in csv.DictReader(stream, delimiter="\t"))
            rows.append([LABELS[d], domains[d], integer(inputs[d]["documents"]),
                         integer(inputs[d]["queries"]), integer(count)])
        note(doc, f"Bảng {part}.1. Quy mô dữ liệu riêng của Phần {part}.")
        table(doc, ["Dataset", "Miền / tác vụ", "Document", "Query test", "Qrels test"],
              rows, [2.8, 5.8, 2.5, 2.5, 3.0], font=10.5)
        paragraph(doc, f"Tổng theo từng bộ của Phần {part}: {integer(sum(inputs[d]['documents'] for d in datasets))} document, {integer(sum(inputs[d]['queries'] for d in datasets))} query test. Corpus được giữ riêng cho từng dataset; không gộp corpus của các bộ.")
        if part == "I":
            paragraph(doc, "Dữ liệu sử dụng định dạng BEIR có sẵn. Corpus là tập tài liệu ứng viên; query là đầu vào tìm kiếm; qrels cung cấp nhãn liên quan để đánh giá. Các số lượng trong bảng được đếm từ file local thực tế.")
            paragraph(doc, "ArguAna có 5 document ID trong qrels không có trong corpus local. Query và qrels gốc vẫn được giữ; các positive thiếu không thể được truy hồi và vẫn ảnh hưởng mẫu số liên quan. NFCorpus và TREC-COVID có nhiều positive/query, nên Recall@100 cần được diễn giải cùng đặc điểm bộ nhãn.")
            paragraph(doc, "Các phiên BM25 cũ từng bỏ document có contents rỗng: SCIDOCS 344, FiQA 38 và TREC-COVID 42.140. Generator đã được sửa để giữ document gốc, kể cả title-only; kết quả trong phần này dùng các lượt chạy thành công đã kiểm tra đủ số document.")
        else:
            doc.add_page_break()
            doc.add_heading(f"{part}.2. Chuẩn bị và chuẩn hóa dữ liệu mở rộng", level=2)
            append_blocks(doc, blocks["2. Chuẩn bị ba dataset tiếng Việt"])

        # Independent data exports and independent figures per group.
        group_dir = output / slug
        group_dir.mkdir(exist_ok=True)
        for key, filename in [("ndcg", "ndcg_at_10.csv"), ("recall", "recall_at_100.csv")]:
            with (group_dir / filename).open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["dataset", *MODELS])
                writer.writerows([[d, *[values[key][d][m] for m in MODELS]] for d in datasets])
            plot_heatmap(values[key], group_dir / (key+"_heatmap.png"),
                         "nDCG@10" if key == "ndcg" else "Recall@100", datasets=datasets)
        group_records = [r for r in records if r["dataset"] in datasets]
        with (group_dir / "sources.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(group_records[0]))
            writer.writeheader()
            writer.writerows(group_records)
        means = chart_means(values["ndcg"], datasets, group_dir / "ndcg_macro.png")
        with (group_dir / "macro_ndcg.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["model", "macro_ndcg_at_10", "dataset_count"])
            writer.writerows([[m, means[m], len(datasets)] for m in MODELS])

        section(doc, True)
        number = 2 if part == "I" else 3
        figure_number = 0
        for index, key in enumerate(["ndcg", "recall"]):
            if index:
                doc.add_page_break()
            metric = "nDCG@10" if key == "ndcg" else "Recall@100"
            doc.add_heading(f"{part}.{number+index}. Kết quả {metric}", level=2)
            paragraph(doc, f"Bảng {part}.{index+2}. Kết quả trên {len(datasets)} dataset của Phần {part}. Chữ đậm và ô xanh biểu thị điểm cao nhất trong tám mô hình trên cùng một dataset; giữ 5 chữ số thập phân.")
            best = {(i, j+1) for i, d in enumerate(datasets) for j, m in enumerate(MODELS)
                    if values[key][d][m] == max(values[key][d].values())}
            table(doc, ["Dataset", *MODELS],
                  [[LABELS[d], *[fmt(values[key][d][m]) for m in MODELS]] for d in datasets],
                  [3.15]+[2.9]*8, font=10.5, best=best)
            figure_number += 1
            figure(doc, group_dir / (key+"_heatmap.png"),
                   f"Hình {part}.{figure_number}. {metric} — chỉ các dataset thuộc Phần {part}. Thang màu 0–1.", 23.5)
            if key == "recall":
                note(doc, "BM25+CE giữ tập ứng viên top-100 của BM25 nên Recall@100 bằng nhau; sự thay đổi thứ tự được phản ánh ở nDCG@10.")
            if add_reference_charts:
                for chart, chart_datasets in reference_bar_charts(values[key], datasets, group_dir, metric, part):
                    doc.add_page_break()
                    doc.add_heading(f"{metric} — So sánh từng dataset", level=2)
                    paragraph(doc, "Các cột biểu diễn điểm thực nghiệm của cùng tám mô hình. BM25 được tô xanh, các mô hình còn lại tô cam; điểm trên cột giữ 5 chữ số thập phân. Các ô dùng cùng thang trục tung để thuận tiện đối chiếu.")
                    figure_number += 1
                    figure(doc, chart,
                           f"Hình {part}.{figure_number}. Biểu đồ cột {metric} trên "
                           + ", ".join(LABELS[d] for d in chart_datasets) + ".", 26.0)

        section(doc)
        summary_number = 4 if part == "I" else 5
        doc.add_heading(f"{part}.{summary_number}. Tổng hợp và nhận xét", level=2)
        paragraph(doc, f"Macro-average dưới đây chỉ tính trên {len(datasets)} dataset thuộc Phần {part}. Mỗi dataset có trọng số bằng nhau; không lấy trung bình chung hai nhóm dữ liệu.")
        table(doc, ["Mô hình", f"nDCG@10 trung bình ({len(datasets)} dataset)"],
              [[m, fmt(means[m])] for m in MODELS], [5.0, 11.6], font=11)
        figure(doc, group_dir / "ndcg_macro.png",
               f"Hình {part}.{figure_number+1}. Trung bình nDCG@10 riêng của Phần {part}; màu xanh ngọc đánh dấu giá trị cao nhất.", 15.8)
        if part == "I":
            paragraph(doc, f"BM25+CE đạt trung bình {fmt(means['BM25+CE'])}, cao hơn BM25 ({fmt(means['BM25'])}) và cải thiện nDCG@10 trên cả sáu dataset. Trong phạm vi tám mô hình, BM25+CE đứng đầu NFCorpus, SciFact, SCIDOCS, FiQA và TREC-COVID.")
            paragraph(doc, f"ArguAna có thứ hạng khác: DeepCT đạt {fmt(values['ndcg']['arguana']['DeepCT'])}, nhỉnh hơn TAS-B ({fmt(values['ndcg']['arguana']['TAS-B'])}) và BM25+CE ({fmt(values['ndcg']['arguana']['BM25+CE'])}). Chưa có kiểm định thống kê nên chênh lệch nhỏ này chỉ được mô tả theo lượt chạy đã ghi.")
            paragraph(doc, "Kết luận Phần I: reranking đem lại lợi ích nhất quán so với BM25 trên sáu bộ được khảo sát, nhưng không đứng đầu mọi tác vụ. Các kết quả là thực nghiệm của project với cấu hình đã nêu, không phải số liệu chép từ bài báo hoặc cam kết tái lập chính xác điểm của tác giả.")
        else:
            paragraph(doc, f"BM25 đạt trung bình {fmt(means['BM25'])}, đứng đầu cả ba dataset. BM25+CE đạt {fmt(means['BM25+CE'])}, đứng thứ hai từng bộ; ColBERT có trung bình {fmt(means['ColBERT'])}. Kết quả chỉ áp dụng cho các checkpoint pretrained và cấu hình đã triển khai.")
            paragraph(doc, "Không quy toàn bộ chênh lệch giữa Phần I và Phần II cho ngôn ngữ: các nhóm đồng thời khác miền, tác vụ và cách tạo qrels. Các mô hình chưa được tối ưu hoặc fine-tune bằng nhãn của ba bộ mở rộng.")
            doc.add_page_break()
            doc.add_heading(f"{part}.6. Phân tích BM25+CE và giới hạn", level=2)
            append_blocks(doc, blocks["8. Vì sao BM25+CE giảm điểm trên tiếng Việt?"])
            paragraph(doc, "Giới hạn riêng của dữ liệu mở rộng: qrels QA chuyển đổi có thể chưa đầy đủ; corpus có passage từ nhiều split; overlap giữa train và test cần được kiểm soát trước huấn luyện supervised. ViQuAD dùng bản mirror có unanswerable, không khẳng định tái lập bản 1.1 gốc. Điểm retrieval không thay thế điểm QA.")
            paragraph(doc, "Kết luận Phần II: BM25 là baseline mạnh nhất trong ba bộ tiếng Việt và tám model đã khảo sát. Việc dùng cùng checkpoint từ thí nghiệm BEIR cho thấy cần kiểm tra khả năng chuyển giao trước khi áp dụng model neural; chưa thể suy rộng thành kết luận về mọi model tiếng Việt/đa ngôn ngữ.")
        section_manifest[slug] = {"datasets": datasets, "models": MODELS, "results": len(group_records)}

    doc.add_page_break()
    doc.add_heading("Tái lập và tài liệu tham khảo", level=2)
    paragraph(doc, "Bản báo cáo này chỉ đổi cách tổ chức, không thay số liệu: Phần I có 48 kết quả, Phần II có 24 kết quả, tổng cộng 72/72 cặp đã kiểm tra. Thời điểm chốt và nguồn artifact được giữ như bản trước. Không có bảng hay biểu đồ kết quả gộp hai nhóm.")
    table(doc, ["Thư mục", "Tệp kết quả và biểu đồ riêng"], [
        ["dataset_bai_bao/", "ndcg_at_10.csv; recall_at_100.csv; macro_ndcg.csv; sources.csv; các hình PNG/PDF cho 6 dataset."],
        ["dataset_mo_rong/", "Các tệp tương ứng, chỉ chứa 3 dataset mở rộng."],
    ], [5.3, 11.3], font=10.5)
    paragraph(doc, "Các lượt chạy gốc nằm tại beir/runs/<model>/results/<run_id>/<dataset>/. File sources.csv trong từng thư mục chỉ rõ nguồn cho từng kết quả. Báo cáo không tự cập nhật theo các lượt chạy về sau; chưa có đánh giá nhiều seed hay kiểm định ý nghĩa thống kê.")
    original_final = blocks["10. Kết luận và khả năng tái lập"]
    references_started = False
    for element in original_final:
        text = "".join(element.xpath(".//w:t/text()"))
        if text == "Tài liệu và nguồn dữ liệu":
            references_started = True
        if references_started:
            doc.element.body.insert(-1, deepcopy(element))

    # Remove old, now unused combined charts from the package relationships.
    used = set(doc.element.xpath(".//@r:embed"))
    for rel_id, relation in list(doc.part.rels.items()):
        if relation.reltype.endswith("/image") and rel_id not in used:
            doc.part.drop_rel(rel_id)
    doc.core_properties.title = "Báo cáo đánh giá — Phần I: dataset BEIR; Phần II: dataset mở rộng"
    path = output / "BAO_CAO_KET_QUA_DANH_GIA_HAI_PHAN.docx"
    doc.save(path)
    save(output / "report_structure.json", {
        "sections": section_manifest, "results": 72, "preserves_original_values": True,
        "reference_style_bar_charts": add_reference_charts,
        "source_word": str(original.resolve()), "source_word_sha256": hashlib.sha256(original.read_bytes()).hexdigest(),
        "report_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    })
    print(path.resolve())
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--add-reference-charts", action="store_true")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve(), args.add_reference_charts)
